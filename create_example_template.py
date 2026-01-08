"""
Create a sample PowerPoint template for testing.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# Create a presentation with some basic styling
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Weekly Project Report"
subtitle.text = "Template Example\nWeek XX, 2024"

# Slide 2: Content Slide
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Project Updates"
tf = content.text_frame
tf.text = "Key Accomplishments:"

p = tf.add_paragraph()
p.text = "Completed feature development"
p.level = 1

p = tf.add_paragraph()
p.text = "Conducted code reviews"
p.level = 1

# Slide 3: Another Content Slide
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Next Steps"
tf = content.text_frame
tf.text = "Upcoming Tasks:"

p = tf.add_paragraph()
p.text = "Deploy to production"
p.level = 1

p = tf.add_paragraph()
p.text = "User acceptance testing"
p.level = 1

# Save the presentation
prs.save('example_template.pptx')
print("✅ Created example_template.pptx")
